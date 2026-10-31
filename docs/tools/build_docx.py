# -*- coding: utf-8 -*-
"""
Сборка книги APEX в DOCX для последующего экспорта в PDF.

Читает собранный Markdown, разбирает его на блоки (заголовки, абзацы,
код, таблицы, списки) и вёрстывает в Word с обложкой, оглавлением,
колонтитулом и нумерацией страниц.
"""

import re
import sys
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = sys.argv[1] if len(sys.argv) > 1 else r"\\wsl.localhost\Ubuntu-24.04\home\nikita\dev\apex\docs\APEX-Книга.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else r"C:\Users\n-kol\AppData\Local\Temp\apex-shop\out\APEX-Книга.docx"
SUBTITLE = sys.argv[3] if len(sys.argv) > 3 else "полная документация проекта"
FOOTER = sys.argv[4] if len(sys.argv) > 4 else "APEX · документация проекта"

# Режим вёрстки: "flow" — главы идут подряд, без разрыва страницы перед каждой.
# Для справочников с короткими главами это обязательно, иначе полстраницы пустует.
FLOW = len(sys.argv) > 5 and sys.argv[5] == "flow"

# Палитра: акцент сайта — фиолетовый
ACCENT = RGBColor(0x6D, 0x28, 0xD9)
ACCENT_DARK = RGBColor(0x4C, 0x1D, 0x95)
TEXT = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
CODE_FILL = "F3F4F6"
HEAD_FILL = "EDE9FE"
BODY_FONT = "Segoe UI"
CODE_FONT = "Consolas"


# ---------------------------------------------------------------- утилиты

def shade(element, fill):
    """Заливка фона для абзаца или ячейки."""
    pr = element.get_or_add_pPr() if element.tag.endswith("}p") else element.get_or_add_tcPr()
    pr.append(parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                        f'w:val="clear" w:color="auto" w:fill="{fill}"/>'))


def run_shade(run, fill):
    run._element.get_or_add_rPr().append(
        parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                  f'w:val="clear" w:color="auto" w:fill="{fill}"/>'))


def left_border(paragraph, color="6D28D9", size=12):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(parse_xml(
        f'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:left w:val="single" w:sz="{size}" w:space="8" w:color="{color}"/></w:pBdr>'))


def add_page_field(paragraph, kind="PAGE"):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), kind)
    run._element.addnext(fld)


def set_font(run, name, size=None, bold=None, color=None, mono=False):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if mono:
        run._element.get_or_add_rPr().append(
            parse_xml('<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                      'w:ascii="Consolas" w:hAnsi="Consolas" w:cs="Consolas"/>'))


# ------------------------------------------------- разбор строки на куски

INLINE = re.compile(
    r"(\*\*.+?\*\*)"      # жирный
    r"|(`[^`]+`)"         # код
    r"|(\[([^\]]+)\]\([^)]*\))"  # ссылка
    r"|(\*[^*]+\*)"       # курсив
)


def add_inline(paragraph, text, size=10.5, color=TEXT, bold_all=False):
    """Раскладывает markdown-строку на форматированные куски."""
    pos = 0

    for match in INLINE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_font(run, BODY_FONT, size, bold_all, color)

        piece = match.group(0)

        if piece.startswith("**"):
            run = paragraph.add_run(piece[2:-2])
            set_font(run, BODY_FONT, size, True, color)
        elif piece.startswith("`"):
            run = paragraph.add_run(piece[1:-1])
            set_font(run, CODE_FONT, size - 1, bold_all, ACCENT_DARK, mono=True)
            run_shade(run, CODE_FILL)
        elif piece.startswith("["):
            run = paragraph.add_run(match.group(4))
            set_font(run, BODY_FONT, size, bold_all, color)
        elif piece.startswith("*"):
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
            set_font(run, BODY_FONT, size, bold_all, color)

        pos = match.end()

    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, BODY_FONT, size, bold_all, color)


# ---------------------------------------------------------------- документ

def build():
    with open(SRC, encoding="utf-8") as handle:
        lines = handle.read().replace("\r\n", "\n").split("\n")

    document = Document()

    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    # Настоящие стили заголовков: из них LibreOffice строит закладки PDF,
    # поэтому по книге можно прыгать через оглавление вьюера
    for name, size, color, before, after in [
        ("Heading 1", 26, ACCENT, 0, 18),
        ("Heading 2", 18, ACCENT_DARK, 0, 12),
        ("Heading 3", 13, TEXT, 14, 6),
    ]:
        style = document.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.font.italic = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = name != "Heading 1"

    # --- обложка ---
    for _ in range(6):
        document.add_paragraph()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("APEX")
    set_font(run, BODY_FONT, 64, True, ACCENT)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(SUBTITLE)
    set_font(run, BODY_FONT, 20, False, TEXT)

    spacer = document.add_paragraph()
    spacer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = spacer.add_run("─" * 22)
    set_font(run, BODY_FONT, 12, False, ACCENT)

    site = document.add_paragraph()
    site.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = site.add_run("mc-apex.ru")
    set_font(run, BODY_FONT, 14, True, ACCENT_DARK)

    for _ in range(9):
        document.add_paragraph()

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("механики · архитектура · код\n{}".format(date.today().strftime("%d.%m.%Y")))
    set_font(run, BODY_FONT, 10, False, MUTED)

    document.add_page_break()

    # --- разбор содержимого ---
    blocks = []
    i = 0
    skip_toc = False

    while i < len(lines):
        line = lines[i]

        # Собственное оглавление строим сами, а титул — обложкой
        if line.startswith("# APEX"):
            i += 1
            continue
        if line.startswith("## Оглавление"):
            skip_toc = True
            i += 1
            continue
        if skip_toc:
            if line.startswith("# "):
                skip_toc = False
            else:
                i += 1
                continue

        if line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("# "):
            blocks.append(("part", line[2:].strip()))
        elif line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(code)))
        elif line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
            rows = [line]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            blocks.append(("table", rows))
            continue
        elif re.match(r"^[-*]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i]):
                items.append(re.sub(r"^[-*]\s+", "", lines[i]))
                i += 1
            blocks.append(("ul", items))
            continue
        elif re.match(r"^\d+\.\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i]):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i]))
                i += 1
            blocks.append(("ol", items))
            continue
        elif line.strip() == "---":
            blocks.append(("hr", None))
        elif line.strip() == "":
            pass
        else:
            blocks.append(("p", line.strip()))

        i += 1

    # --- оглавление ---
    heading = document.add_paragraph()
    run = heading.add_run("Оглавление")
    set_font(run, BODY_FONT, 24, True, ACCENT)
    heading.paragraph_format.space_after = Pt(14)

    for kind, value in blocks:
        # В оглавлении обратные кавычки лишние — это разметка, не текст
        plain = value.replace("`", "") if isinstance(value, str) else ""

        if kind == "part":
            par = document.add_paragraph()
            par.paragraph_format.space_before = Pt(10)
            run = par.add_run(plain)
            set_font(run, BODY_FONT, 11, True, ACCENT)
        elif kind == "h2":
            numbered = re.match(r"^\d+\.\s", plain) is not None
            par = document.add_paragraph()
            par.paragraph_format.left_indent = Cm(1.3 if numbered else 0.7)
            run = par.add_run(plain)
            set_font(run, BODY_FONT, 10, False, TEXT)

    document.add_page_break()

    # --- содержимое ---
    prev_was_part = False

    for kind, value in blocks:
        if kind == "part":
            par = document.add_heading(value, level=1)
            par.paragraph_format.page_break_before = True
            prev_was_part = True

        elif kind == "h2":
            par = document.add_heading(value, level=2)

            if FLOW:
                par.paragraph_format.space_before = Pt(20)
            elif not prev_was_part:
                # Первая глава части идёт на той же странице, что и её
                # заголовок: иначе заголовок части остаётся один
                par.paragraph_format.page_break_before = True

            prev_was_part = False

        elif kind == "h3":
            par = document.add_heading(value, level=3)

        elif kind == "p":
            par = document.add_paragraph()
            add_inline(par, value)

        elif kind == "code":
            par = document.add_paragraph()
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(8)
            par.paragraph_format.left_indent = Cm(0.3)
            par.paragraph_format.line_spacing = 1.0
            shade(par._p, CODE_FILL)
            left_border(par)
            lines_code = value.split("\n")
            for index, chunk in enumerate(lines_code):
                run = par.add_run(chunk)
                set_font(run, CODE_FONT, 8.5, False, TEXT, mono=True)
                if index < len(lines_code) - 1:
                    run.add_break()

        elif kind == "ul" or kind == "ol":
            for index, item in enumerate(value, start=1):
                par = document.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.8)
                par.paragraph_format.space_after = Pt(3)
                marker = "•  " if kind == "ul" else f"{index}.  "
                run = par.add_run(marker)
                set_font(run, BODY_FONT, 10.5, True, ACCENT)
                add_inline(par, item)

        elif kind == "table":
            grid = []
            for row in value:
                cells = [cell.strip() for cell in row.strip().strip("|").split("|")]
                grid.append(cells)

            width = max(len(row) for row in grid)
            table = document.add_table(rows=len(grid), cols=width)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            for row_index, cells in enumerate(grid):
                for col_index in range(width):
                    cell = table.cell(row_index, col_index)
                    cell.paragraphs[0].text = ""
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    par.paragraph_format.space_before = Pt(2)

                    text = cells[col_index] if col_index < len(cells) else ""
                    add_inline(par, text, size=9, bold_all=(row_index == 0))

                    if row_index == 0:
                        shade(cell._tc, HEAD_FILL)

            document.add_paragraph().paragraph_format.space_after = Pt(4)

        elif kind == "hr":
            par = document.add_paragraph()
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
    # --- колонтитул с номером страницы ---
    footer = document.sections[0].footer
    fpar = footer.paragraphs[0]
    fpar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fpar.add_run(FOOTER + " · ")
    set_font(run, BODY_FONT, 8, False, MUTED)
    add_page_field(fpar, "PAGE")

    document.save(OUT)
    print("  сохранено:", OUT)
    print("  блоков разобрано:", len(blocks))


if __name__ == "__main__":
    build()
